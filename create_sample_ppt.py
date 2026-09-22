from pptx import Presentation
from pptx.util import Inches, Pt

def create_sample_presentation(filename="sample_study_presentation.pptx"):
    prs = Presentation()

    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]

    title1.text = "Introduction to Machine Learning & Neural Networks"
    subtitle1.text = "Computer Science Study Guide 2026"

    # Speaker Notes for Slide 1
    notes_slide1 = slide1.notes_slide
    text_frame1 = notes_slide1.notes_text_frame
    text_frame1.text = "WELCOME NOTES: Emphasize that Machine Learning is a subset of AI. Remind students that the midterm exam will cover Supervised vs Unsupervised learning definitions!"

    # Slide 2: Supervised Learning
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    shapes2 = slide2.shapes
    title2 = shapes2.title
    body2 = shapes2.placeholders[1]

    title2.text = "Supervised Learning Fundamentals"
    tf2 = body2.text_frame
    tf2.text = "Supervised learning relies on labeled training data."
    
    p = tf2.add_paragraph()
    p.text = "Key Algorithms:"
    p.level = 1

    p2 = tf2.add_paragraph()
    p2.text = "Linear Regression (predicting continuous values)"
    p2.level = 2

    p3 = tf2.add_paragraph()
    p3.text = "Decision Trees & Random Forests (classification)"
    p3.level = 2

    # Speaker Notes for Slide 2
    notes_slide2 = slide2.notes_slide
    text_frame2 = notes_slide2.notes_text_frame
    text_frame2.text = "IMPORTANT PRESENTER NOTE: Make sure to explain the difference between regression (predicting temperature/price) and classification (predicting spam/not spam). Students often confuse these two on the quiz."

    # Slide 3: Table of Comparison
    slide3 = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    title_box = slide3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    title_box.text_frame.text = "Comparison: Supervised vs Unsupervised Learning"

    rows, cols = 3, 3
    left, top, width, height = Inches(0.5), Inches(1.8), Inches(9), Inches(3)
    table = slide3.shapes.add_table(rows, cols, left, top, width, height).table

    table.cell(0, 0).text = "Feature"
    table.cell(0, 1).text = "Supervised Learning"
    table.cell(0, 2).text = "Unsupervised Learning"

    table.cell(1, 0).text = "Data Type"
    table.cell(1, 1).text = "Labeled Data"
    table.cell(1, 2).text = "Unlabeled Data"

    table.cell(2, 0).text = "Goal"
    table.cell(2, 1).text = "Predict outcomes"
    table.cell(2, 2).text = "Find hidden patterns / clusters"

    # Speaker Notes for Slide 3
    notes_slide3 = slide3.notes_slide
    text_frame3 = notes_slide3.notes_text_frame
    text_frame3.text = "SPEAKER NOTE UNDER SLIDE 3: Clustering (k-means) and Principal Component Analysis (PCA) are prime examples of Unsupervised Learning."

    prs.save(filename)
    print(f"Sample presentation created successfully: {filename}")

if __name__ == "__main__":
    create_sample_presentation()
